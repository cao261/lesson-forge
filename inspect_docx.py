import sys
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
ep_dir = Path("workspace/series_line_circle/ep119")
teacher_docx = ep_dir / "【P119精讲】直线的倾斜角与斜率_教师备课版.docx"
student_docx = ep_dir / "【P119精讲】直线的倾斜角与斜率_学生随堂版.docx"

print("=== Checking Teacher Docx ===")
doc_t = Document(teacher_docx)
print(f"Paragraphs: {len(doc_t.paragraphs)}, Tables: {len(doc_t.tables)}")
t_text = "\n".join(p.text for p in doc_t.paragraphs if p.text)
for ex_id in ["ex1", "ex2", "ex3", "ex4", "ex5", "ex6", "ex7", "ex8", "ex9", "ex10"]:
    title_hit = f"抢答判断题" in t_text or f"例题" in t_text or f"综合练习" in t_text
print(f"Teacher doc text sample:\n{t_text[:500]}...")

print("\n=== Checking Student Docx ===")
doc_s = Document(student_docx)
print(f"Paragraphs: {len(doc_s.paragraphs)}, Tables: {len(doc_s.tables)}")
s_text = "\n".join(p.text for p in doc_s.paragraphs if p.text)
print(f"Student doc text sample:\n{s_text[:500]}...")

# Check inline shapes (images embedded)
print("\n=== Checking Images in Docs ===")
print("Teacher doc inline shapes count:", len(doc_t.inline_shapes))
print("Student doc inline shapes count:", len(doc_s.inline_shapes))